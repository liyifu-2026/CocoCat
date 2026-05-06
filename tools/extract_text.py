#!/usr/bin/env python3
"""Extract text from various file formats and output as JSON to stdout."""
import json
import os
import sys


def extract_text(filepath: str) -> dict:
    ext = os.path.splitext(filepath)[1].lower()

    # --- PDF ---
    if ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(filepath)
            pages = []
            for page in doc:
                pages.append({"page": page.number + 1, "text": page.get_text()})
            text = "\n\n".join(p["text"] for p in pages)
            return {"text": text, "pages": len(pages)}
        except Exception:
            pass  # fall through to text fallback

    # --- Word ---
    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(filepath)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n\n".join(paragraphs)
            return {"text": text, "pages": 1}
        except Exception:
            pass

    # --- Excel ---
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(filepath, read_only=True, data_only=True)
            rows = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                sheet_rows = []
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v) for v in row if v is not None]
                    if vals:
                        sheet_rows.append(" | ".join(vals))
                if sheet_rows:
                    rows.append(f"### Sheet: {sheet}\n" + "\n".join(sheet_rows))
            text = "\n\n".join(rows)
            return {"text": text, "pages": len(wb.sheetnames)}
        except Exception:
            pass

    # --- PowerPoint ---
    if ext == ".pptx":
        try:
            from pptx import Presentation
            prs = Presentation(filepath)
            slides = []
            for i, slide in enumerate(prs.slides, 1):
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        texts.append(shape.text)
                if texts:
                    slides.append(f"### Slide {i}\n" + "\n".join(texts))
            text = "\n\n".join(slides)
            return {"text": text, "pages": len(prs.slides)}
        except Exception:
            pass

    # --- HTML ---
    if ext in (".htm", ".html"):
        try:
            from bs4 import BeautifulSoup
            with open(filepath, "r", encoding="utf-8", errors="replace") as f:
                soup = BeautifulSoup(f, "html.parser")
            for tag in soup(["script", "style", "nav", "footer", "header"]):
                tag.decompose()
            text = soup.get_text(separator="\n")
            lines = [l.strip() for l in text.split("\n") if l.strip()]
            text = "\n".join(lines)
            return {"text": text, "pages": 1}
        except Exception:
            pass

    # --- Text fallback ---
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    return {"text": text, "pages": 1}


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"error": "Usage: extract_text.py <filepath>"}))
        sys.exit(1)
    try:
        result = extract_text(sys.argv[1])
        print(json.dumps(result))
    except Exception as e:
        print(json.dumps({"error": str(e)}))
        sys.exit(1)
